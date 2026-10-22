#!/usr/bin/env python3
"""Render a cover letter written in Markdown into a signed PDF and an editable DOCX.

Usage:
    python3 render_cover_letter.py LETTER.md OUT_DIR [--roster AUTHORS.yml] [--assets DIR]
                                   [--max-pages N] [--check-docx]

LETTER.md has YAML front matter (keys in ../assets/cover_letter.md) followed by the letter body,
from "Dear Editor," to the declarations. The closing, signer label and signer blocks are generated.

Outputs (fixed names, see ../references/output-layout.md):
    OUT_DIR/cover_letter.pdf            XeLaTeX, layout from ../assets/cover_letter_layout.tex
    OUT_DIR/cover_letter_editable.docx  same text and layout for Word editing
Prints a JSON report: pages, signers, signers without a signature image, LaTeX warnings.
Exit code: 0 = rendered within --max-pages; 1 = the PDF (or, with --check-docx, the LibreOffice
rendering of the DOCX) exceeds --max-pages, or a tool failed; 2 = bad input.

Signers: each entry of `signers` is an author id, or a mapping with `id` plus any of name, title,
affiliation (list of lines) and email. Facts for an id come from the roster YAML (--roster, format in
../assets/authors.example.yml); keys in the letter override the roster.
Images: logos are looked up by file stem in <assets>/logos/, signatures as <id>.* in <assets>/signatures/.
<assets> is --assets, else the roster's folder, else the letter's folder. A signer without a signature
image keeps an empty space of the same height, so signed and unsigned signers can share one letter.

Requires: pandoc, latexmk + XeLaTeX with the TeX Gyre Termes font, pdfinfo (poppler-utils), Python
packages python-docx and PyYAML; LibreOffice (soffice) only for --check-docx.
Temporary build files are created inside OUT_DIR and removed afterwards.
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.styles.style import CharacterStyle, ParagraphStyle

SKILL = Path(__file__).resolve().parents[1]
LAYOUT = SKILL / "assets/cover_letter_layout.tex"
PDF_NAME, DOCX_NAME = "cover_letter.pdf", "cover_letter_editable.docx"
# Same page geometry as the LaTeX layout.
PAGE = {"page_width": Cm(21.0), "page_height": Cm(29.7), "top_margin": Cm(1.2), "bottom_margin": Cm(1.2),
        "left_margin": Cm(1.6), "right_margin": Cm(1.6)}
TEXT_WIDTH = Cm(21.0 - 2 * 1.6)
FONT, SIZE = "Times New Roman", Pt(10)


class InputError(Exception):
    """Bad letter, roster or asset; reported with exit code 2."""


def load_letter(path: Path) -> tuple[dict, str]:
    match = re.match(r"^---\n(.*?)\n---\n(.*)$", path.read_text(encoding="utf-8"), re.S)
    if not match:
        raise InputError(f"{path} needs YAML front matter between --- lines")
    meta = yaml.safe_load(match.group(1)) or {}
    for key in ("title", "authors", "journal", "signers"):
        if not meta.get(key):
            raise InputError(f"front matter key '{key}' is missing")
    return meta, match.group(2).strip()


def find_asset(folder: Path, stem: str) -> Path | None:
    hits = sorted(p for p in folder.glob(f"{stem}.*") if p.is_file()) if folder.is_dir() else []
    return hits[0].resolve() if hits else None


def resolve_signers(entries: list, roster: dict, signatures: Path) -> list[dict]:
    people = (roster or {}).get("people", {})
    signers = []
    for entry in entries:
        entry = {"id": entry} if isinstance(entry, str) else dict(entry)
        if not entry.get("id"):
            raise InputError(f"signer {entry} needs an 'id'")
        person = {**people.get(entry["id"], {}), **entry}
        affiliation = person.get("affiliation") or []
        signer = {
            "id": entry["id"],
            "name": person.get("name"),
            "title": person.get("title"),
            "affiliation": [affiliation] if isinstance(affiliation, str) else list(affiliation),
            "email": person.get("email"),
            "signature": find_asset(signatures, entry["id"]),
        }
        missing = [k for k in ("name", "title", "email") if not signer[k]]
        if missing:
            raise InputError(f"signer {entry['id']}: {missing} missing in roster and front matter")
        signers.append(signer)
    return signers


def convert_markdown(text: str, inline: bool = False) -> str:
    latex = subprocess.run(["pandoc", "-f", "markdown", "-t", "latex", "--wrap=none"], input=text, text=True,
                           capture_output=True, check=True).stdout.strip()
    return latex.removeprefix("\\par").strip() if inline else latex


def signer_label(meta: dict, signers: list[dict]) -> str:
    return meta.get("signers_label", "Corresponding Author" + ("s:" if len(signers) > 1 else ":"))


def render_pdf(meta: dict, body: str, signers: list[dict], logo_files: list[Path], out: Path) -> tuple[int, list[str]]:
    width = {1: 0.45, 2: 0.47}.get(len(signers), 0.315)
    blocks = []
    for i, s in enumerate(signers):
        image = rf"\includegraphics[height=0.85cm]{{{s['signature']}}}" if s["signature"] else ""
        lines = r"\\ ".join([convert_markdown(s["title"], True)] + [convert_markdown(x, True) for x in s["affiliation"]])
        blocks.append(rf"\signer{{{image}}}{{{convert_markdown(s['name'], True)}}}{{{lines}}}{{{s['email']}}}")
        if i % 3 == 2 and i + 1 < len(signers):
            blocks.append(r"\\[8pt]")
        elif i + 1 < len(signers):
            blocks.append(r"\hfill")
    fill = {
        "SIGNER_WIDTH": str(width),
        "LOGOS": r"\hfill".join(rf"\includegraphics[height=0.95cm]{{{p}}}" for p in logo_files),
        "TITLE": convert_markdown(meta["title"], True),
        "AUTHORS": convert_markdown(", ".join(meta["authors"]), True),
        "DATE": f"\n\\vspace{{2pt}}\n{meta['date']}\n" if meta.get("date") else "",
        "BODY": convert_markdown(body),
        "CLOSING": meta.get("closing", "Sincerely,"),
        "SIGNER_LABEL": signer_label(meta, signers),
        "SIGNERS": "\n".join(blocks),
    }
    tex = LAYOUT.read_text(encoding="utf-8")
    for key, value in fill.items():
        tex = tex.replace(f"<<{key}>>", value)
    with tempfile.TemporaryDirectory(dir=out, prefix=".render_") as tmp:
        (Path(tmp) / "cover_letter.tex").write_text(tex, encoding="utf-8")
        run = subprocess.run(["latexmk", "-xelatex", "-interaction=nonstopmode", "-halt-on-error", "cover_letter.tex"],
                             cwd=tmp, capture_output=True, text=True)
        log_file = Path(tmp) / "cover_letter.log"
        log = log_file.read_text(errors="replace") if log_file.exists() else ""
        if run.returncode != 0:
            raise RuntimeError("XeLaTeX build failed:\n" + run.stdout[-3000:])
        (out / PDF_NAME).write_bytes((Path(tmp) / "cover_letter.pdf").read_bytes())
    warnings = sorted(set(re.findall(r"^(?:Overfull|Underfull) \\hbox.*$|^.*Missing character.*$", log, re.M)))
    return count_pages(out / PDF_NAME), warnings


def count_pages(pdf: Path) -> int:
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, check=True).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.M)
    if not match:
        raise RuntimeError(f"pdfinfo gave no page count for {pdf}")
    return int(match.group(1))


def apply_font(style, size=None) -> None:
    style.font.name = FONT
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        style.font.size = size


def add_run(paragraph, text: str, bold: bool = False, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if size is not None:
        run.font.size = size
    return run


def write_docx(meta: dict, body: str, signers: list[dict], logo_files: list[Path], out: Path) -> None:
    path = out / DOCX_NAME
    subprocess.run(["pandoc", "-f", "markdown", "-t", "docx", "-o", str(path)], input=body, text=True,
                   capture_output=True, check=True)
    doc = Document(str(path))
    for name, value in PAGE.items():
        setattr(doc.sections[0], name, value)
    for style in doc.styles:
        if isinstance(style, (ParagraphStyle, CharacterStyle)):
            apply_font(style, SIZE if style.name in ("Normal", "Body Text", "First Paragraph", "Compact") else None)
        if isinstance(style, ParagraphStyle):
            style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(0), Pt(3)
    first = doc.paragraphs[0]
    if logo_files:
        logos = first.insert_paragraph_before()
        stops = logos.paragraph_format.tab_stops
        if len(logo_files) == 3:
            stops.add_tab_stop(int(TEXT_WIDTH / 2), WD_TAB_ALIGNMENT.CENTER)
        if len(logo_files) > 1:
            stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
        for i, logo in enumerate(logo_files):
            if i:
                logos.add_run("\t")
            logos.add_run().add_picture(str(logo), height=Cm(0.95))
    heading = first.insert_paragraph_before()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(heading, "Cover Letter", True, Pt(12))
    head = first.insert_paragraph_before()
    add_run(head, "Title: ", True)
    add_run(head, meta["title"])
    head.add_run().add_break()
    add_run(head, "Authors: ", True)
    add_run(head, ", ".join(meta["authors"]))
    if meta.get("date"):
        first.insert_paragraph_before(str(meta["date"]))
    doc.add_paragraph(meta.get("closing", "Sincerely,"))
    add_run(doc.add_paragraph(), signer_label(meta, signers), True)
    table = doc.add_table(rows=(len(signers) + 2) // 3, cols=min(len(signers), 3))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, s in enumerate(signers):
        cell = table.rows[i // 3].cells[i % 3]
        top = cell.paragraphs[0]
        top.paragraph_format.line_spacing = Cm(0.95)
        if s["signature"]:
            top.add_run().add_picture(str(s["signature"]), height=Cm(0.85))
        for text, bold in [(s["name"], True), (s["title"], False), *[(x, False) for x in s["affiliation"]],
                           (f"Email: {s['email']}", False)]:
            add_run(cell.add_paragraph(), text, bold, Pt(9))
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
    # No author name in file properties (double-blind journals check them).
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(str(path))


def count_docx_pages(docx: Path) -> int:
    with tempfile.TemporaryDirectory(dir=docx.parent, prefix=".render_") as tmp:
        profile = Path(tmp) / "lo-profile"
        subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                        "--outdir", tmp, str(docx)], capture_output=True, check=True)
        return count_pages(Path(tmp) / f"{docx.stem}.pdf")


def main() -> int:
    parser = argparse.ArgumentParser(description="Render a Markdown cover letter into a signed PDF and an editable DOCX.")
    parser.add_argument("letter", type=Path)
    parser.add_argument("out_dir", type=Path)
    parser.add_argument("--roster", type=Path, help="author roster YAML (see assets/authors.example.yml)")
    parser.add_argument("--assets", type=Path, help="folder with logos/ and signatures/ (default: roster or letter folder)")
    parser.add_argument("--max-pages", type=int, default=1)
    parser.add_argument("--check-docx", action="store_true", help="also count DOCX pages through LibreOffice")
    args = parser.parse_args()
    try:
        if not args.letter.is_file():
            raise InputError(f"{args.letter} not found")
        if args.roster and not args.roster.is_file():
            raise InputError(f"roster {args.roster} not found")
        meta, body = load_letter(args.letter)
        roster = yaml.safe_load(args.roster.read_text(encoding="utf-8")) if args.roster else {}
        assets = (args.assets or (args.roster.parent if args.roster else args.letter.parent)).resolve()
        signers = resolve_signers(meta["signers"], roster, assets / "signatures")
        logo_files = [find_asset(assets / "logos", stem) for stem in meta.get("logos") or []]
        if not all(logo_files):
            missing = [s for s, f in zip(meta["logos"], logo_files) if not f]
            raise InputError(f"logo not found in {assets / 'logos'}: {missing}")
    except InputError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 2
    args.out_dir.mkdir(parents=True, exist_ok=True)
    out = args.out_dir.resolve()
    try:
        pages, warnings = render_pdf(meta, body, signers, logo_files, out)
        write_docx(meta, body, signers, logo_files, out)
        report = {"pdf": str(out / PDF_NAME), "docx": str(out / DOCX_NAME), "pdf_pages": pages,
                  "signers": [s["name"] for s in signers],
                  "signers_without_signature": [s["name"] for s in signers if not s["signature"]],
                  "latex_warnings": warnings}
        if args.check_docx:
            report["docx_pages"] = count_docx_pages(out / DOCX_NAME)
    except (RuntimeError, subprocess.CalledProcessError) as e:
        detail = e.stderr if isinstance(e, subprocess.CalledProcessError) else str(e)
        print(f"ERROR: {detail}", file=sys.stderr)
        return 1
    print(json.dumps(report, indent=2, ensure_ascii=False))
    too_long = pages > args.max_pages or report.get("docx_pages", 0) > args.max_pages
    return 1 if too_long else 0


if __name__ == "__main__":
    sys.exit(main())
